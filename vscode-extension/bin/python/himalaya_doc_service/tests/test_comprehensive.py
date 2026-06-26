"""Comprehensive integration test for all Himalaya doc service tools."""

import sys, os, json, tempfile, zipfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'src'))

from himalaya_doc_service.server import McpServer
from himalaya_doc_service.protocol import tool_result
from himalaya_doc_service.spec import (
    AssetPlan,
    BlockType,
    ChartKind,
    DocumentSpec,
    DocumentType,
    FormulaCell,
    GenerationContract,
    SheetSpec,
    SourceRef,
    SpecBlock,
)
from himalaya_doc_service.themes import load_theme, THEMES, hex_to_rgb
from himalaya_doc_service.quality import assess_document
from himalaya_doc_service.tools import TOOL_REGISTRY, TOOL_SCHEMAS
from pptx import Presentation


def test_imports():
    assert len(TOOL_REGISTRY) == 8
    assert len(THEMES) == 6
    print("  PASS: imports")


def test_comprehensive_spec():
    spec = DocumentSpec(
        title='Full Test', subtitle='All block types',
        author='Himalaya', language='zh-CN',
        theme={'name': 'ocean'},
        blocks=[
            SpecBlock(type=BlockType.HEADING, level=1, text='Intro'),
            SpecBlock(type=BlockType.PARAGRAPH, text='Para'),
            SpecBlock(type=BlockType.BULLETS, items=['A', 'B']),
            SpecBlock(type=BlockType.TABLE, caption='Tbl', headers=['H1'], rows=[['v1']]),
            SpecBlock(type=BlockType.CHART, title='Chart', kind=ChartKind.BAR,
                     labels=['A'], series=[{'name': 'S1', 'values': [1.0]}]),
            SpecBlock(type=BlockType.FORMULA, latex=r'E=mc^2'),
            SpecBlock(type=BlockType.QUOTE, text='Quote', attribution='Author'),
            SpecBlock(type=BlockType.CODE, language='python', code='print(1)'),
            SpecBlock(type=BlockType.TWO_COLUMN,
                     left=[SpecBlock(type=BlockType.PARAGRAPH, text='Left')],
                     right=[SpecBlock(type=BlockType.PARAGRAPH, text='Right')]),
            SpecBlock(type=BlockType.PAGE_BREAK),
            SpecBlock(type=BlockType.HEADING, level=1, text='Final'),
            SpecBlock(type=BlockType.IMAGE, path='/tmp/test_phase0_img.png', alt='img', caption='cap'),
        ],
        sheets=[SheetSpec(name='S1', rows=[['A']], formulas=[FormulaCell(cell='A2', formula='=SUM(A1)')])],
    )
    sd = spec.model_dump()
    assert len(sd['blocks']) == 12
    assert len(sd['sheets']) == 1
    print("  PASS: comprehensive spec")


def test_all_formats():
    spec = DocumentSpec(
        title='Test', blocks=[
            SpecBlock(type=BlockType.HEADING, level=1, text='S1'),
            SpecBlock(type=BlockType.PARAGRAPH, text='Hello'),
            SpecBlock(type=BlockType.TABLE, caption='Data', headers=['A', 'B'], rows=[['1', '2']]),
            SpecBlock(type=BlockType.CHART, title='Sales', kind=ChartKind.BAR,
                     labels=['Q1', 'Q2'], series=[{'name': '2024', 'values': [10, 20]}]),
            SpecBlock(type=BlockType.FORMULA, latex=r'E=mc^2'),
        ],
    )
    sd = spec.model_dump()
    for fmt in ['pptx', 'docx', 'xlsx', 'pdf']:
        out = f'/tmp/test_comp_{fmt}.{fmt}'
        r = TOOL_REGISTRY[f'generate_{fmt}']({'path': out, 'document_spec': sd})
        assert os.path.exists(r['path']), f"{fmt} file not created"
        assert 'quality' in r, f"{fmt} missing quality"
        assert 'manifestPath' in r, f"{fmt} missing manifestPath"
        print(f"  PASS: {fmt.upper()} generation ({os.path.getsize(r['path'])}b)")


def test_chart_image():
    r = TOOL_REGISTRY['generate_chart_image']({
        'output_path': '/tmp/test_chart.png',
        'chart_type': 'bar', 'title': 'Test', 'labels': ['A'], 'series': [{'name': 'X', 'values': [1]}],
    })
    assert os.path.exists(r['path'])
    assert r['size_bytes'] > 100
    print(f"  PASS: chart image ({r['size_bytes']}b)")


def test_formula_image():
    r = TOOL_REGISTRY['render_formula_image']({
        'latex': r'\frac{a}{b}',
        'output_path': '/tmp/test_formula.png',
    })
    assert os.path.exists(r['path'])
    assert r['size_bytes'] > 100
    assert r['rendered_as'] in ('mathtext', 'plain_latex_image')
    print(f"  PASS: formula image ({r['size_bytes']}b)")


def test_pptx_embeds_real_objects_and_office_math():
    from PIL import Image

    img_path = '/tmp/test_pptx_real_image.png'
    Image.new('RGB', (360, 200), '#EEF5FF').save(img_path)
    spec = DocumentSpec(
        title='Real PPTX objects',
        generation_contract=GenerationContract(require_editable_formulas=True),
        blocks=[
            SpecBlock(type=BlockType.HEADING, level=1, text='Objects'),
            SpecBlock(type=BlockType.IMAGE, path=img_path, caption='Source figure'),
            SpecBlock(type=BlockType.TABLE, headers=['Metric', 'Value'], rows=[['A', '1']]),
            SpecBlock(type=BlockType.FORMULA, latex=r'R=\frac{1}{N}\sum_i r_i', equation_number='Eq. (1)'),
        ],
    )
    out = '/tmp/test_pptx_real_objects.pptx'
    r = TOOL_REGISTRY['generate_pptx']({'path': out, 'document_spec': spec.model_dump()})
    with zipfile.ZipFile(out) as zf:
        slide_xml = b''.join(
            zf.read(name)
            for name in zf.namelist()
            if name.startswith('ppt/slides/slide') and name.endswith('.xml')
        )
    assert b'name="Spec Image"' in slide_xml
    assert b'<a:tbl' in slide_xml
    assert b'<m:oMath' in slide_xml
    assert r['quality']['embedded_image_count'] >= 1
    assert r['quality']['rendered_table_count'] >= 1
    assert r['quality']['editable_formula_count'] >= 1
    assert 'formula.editable_native_unavailable' not in '\n'.join(r['quality']['warnings'] + r['quality']['failures'])
    print("  PASS: PPTX embeds image/table/native Office Math")


def test_pptx_asset_manifest_hydrates_placeholders():
    from PIL import Image

    img_path = '/tmp/test_manifest_hydrate_image.png'
    Image.new('RGB', (320, 180), '#FFFFFF').save(img_path)
    manifest = {
        'suggested_image_blocks': [{
            'type': 'image',
            'path': img_path,
            'asset_role': 'figure',
            'caption': 'Extracted source figure',
            'source_path': 'paper.pdf',
            'crop_box': [10, 20, 300, 180],
        }],
        'suggested_table_blocks': [{
            'type': 'table',
            'caption': 'Extracted source table',
            'headers': ['Parameter', 'Value'],
            'rows': [['N', '1000']],
        }],
        'suggested_formula_blocks': [{
            'type': 'formula',
            'latex': r'\alpha_t=\frac{x_t}{N}',
            'metadata': {'render_as': 'office_math'},
        }],
    }
    spec = DocumentSpec(
        title='Hydrate placeholders',
        generation_contract=GenerationContract(require_editable_formulas=True),
        blocks=[
            SpecBlock(type=BlockType.HEADING, level=1, text='Hydrated'),
            SpecBlock(type=BlockType.IMAGE, asset_role='figure', caption='placeholder'),
            SpecBlock(type=BlockType.TABLE, caption='placeholder'),
            SpecBlock(type=BlockType.FORMULA),
        ],
    )
    out = '/tmp/test_manifest_hydrate.pptx'
    r = TOOL_REGISTRY['generate_pptx']({
        'path': out,
        'document_spec': spec.model_dump(),
        'asset_manifest': manifest,
    })
    assert r['quality']['embedded_image_count'] >= 1
    assert r['quality']['rendered_table_count'] >= 1
    assert r['quality']['editable_formula_count'] >= 1
    print("  PASS: asset manifest hydrates PPTX placeholders")


def test_pdf_asset_extraction_batches_and_captioned_figures():
    import fitz

    pdf_path = '/tmp/test_captioned_assets.pdf'
    out_dir = tempfile.mkdtemp(prefix='himalaya-assets-')
    doc = fitz.open()
    for page_no in range(3):
        page = doc.new_page(width=420, height=300)
        page.insert_text((40, 40), f'第{page_no + 1}页 论文内容', fontsize=12)
        if page_no == 0:
            shape = page.new_shape()
            shape.draw_rect(fitz.Rect(80, 80, 340, 185))
            shape.finish(color=(0, 0, 0.8), fill=(0.88, 0.94, 1.0), width=1)
            shape.commit()
            page.insert_text((110, 205), 'Fig. 1.1 UAV swarm reconnaissance framework', fontsize=11)
            page.insert_text((55, 245), r'R=\frac{1}{N}\sum_i r_i', fontsize=11)
    doc.save(pdf_path)
    doc.close()

    r = TOOL_REGISTRY['extract_document_assets']({
        'path': pdf_path,
        'output_dir': out_dir,
        'page_batch_size': 1,
        'resume': False,
        'extract_images': False,
        'extract_tables': False,
        'extract_formula_candidates': True,
        'min_height': 8,
    })

    assert r['pages_scanned'] == 1
    assert r['complete'] is False
    assert r['next_start_page'] == 2
    assert r['asset_count'] >= 1
    assert any(a.get('extraction_method') == 'caption_crop' for a in r['assets'])
    assert any((a.get('role') == 'formula_candidate') for a in r['assets'])
    assert r.get('suggested_formula_blocks')
    assert os.path.exists(r['manifest_path'])
    print(f"  PASS: PDF assets batch + caption crop ({r['asset_count']} assets)")


def test_pdf_asset_extraction_detects_vector_tables():
    import contextlib
    import fitz
    import io

    pdf_path = '/tmp/test_vector_table_assets.pdf'
    out_dir = tempfile.mkdtemp(prefix='himalaya-table-assets-')
    doc = fitz.open()
    page = doc.new_page(width=420, height=300)
    shape = page.new_shape()
    x0, y0, x1, y1 = 60, 80, 360, 180
    for x in [x0, 160, 260, x1]:
        shape.draw_line((x, y0), (x, y1))
    for y in [y0, 115, 150, y1]:
        shape.draw_line((x0, y), (x1, y))
    shape.finish(color=(0, 0, 0), width=1)
    shape.commit()
    for idx, text in enumerate(['Metric', 'Baseline', 'Ours']):
        page.insert_text((75 + idx * 100, 103), text, fontsize=10)
    for row in range(2):
        for col in range(3):
            page.insert_text((75 + col * 100, 138 + row * 35), str(row * 3 + col), fontsize=10)
    page.insert_text((95, 205), 'Table 1. Robustness comparison', fontsize=11)
    doc.save(pdf_path)
    doc.close()

    stdout = io.StringIO()
    with contextlib.redirect_stdout(stdout):
        r = TOOL_REGISTRY['extract_document_assets']({
            'path': pdf_path,
            'output_dir': out_dir,
            'page_batch_size': 1,
            'resume': False,
            'extract_images': False,
            'extract_captioned_figures': False,
            'extract_formula_candidates': False,
            'extract_tables': True,
            'min_height': 8,
        })

    table_assets = [a for a in r['assets'] if a.get('role') == 'table_image']
    assert table_assets
    assert os.path.exists(table_assets[0]['path'])
    assert table_assets[0].get('caption') == 'Table 1. Robustness comparison'
    assert r.get('suggested_table_blocks')
    assert r['suggested_table_blocks'][0]['headers'][:2] == ['Metric', 'Baseline']
    assert 'pymupdf_layout' not in stdout.getvalue()
    print(f"  PASS: PDF vector table extraction ({len(table_assets)} table asset)")


def test_mcp_protocol():
    server = McpServer()
    # initialize
    r = server._dispatch({'jsonrpc': '2.0', 'id': 1, 'method': 'initialize', 'params': {}})
    assert r['result']['serverInfo']['name'] == 'himalaya-doc-service'
    # list tools
    r = server._dispatch({'jsonrpc': '2.0', 'id': 2, 'method': 'tools/list', 'params': {}})
    assert len(r['result']['tools']) == 8
    # call tool
    r = server._dispatch({
        'jsonrpc': '2.0', 'id': 3, 'method': 'tools/call',
        'params': {'name': 'generate_chart_image',
                   'arguments': {'output_path': '/tmp/test_mcp_c.png', 'chart_type': 'pie',
                                 'labels': ['X'], 'series': [{'name': 'Y', 'values': [100]}]}}
    })
    result_text = r['result']['content'][0]['text']
    assert 'path' in result_text
    print("  PASS: MCP protocol (initialize, list, call)")


def test_quality_warnings():
    spec = DocumentSpec(
        blocks=[
            SpecBlock(type=BlockType.TABLE, caption='Empty', headers=[], rows=[]),
            SpecBlock(type=BlockType.FORMULA, latex=''),
            SpecBlock(type=BlockType.CHART, title='Irregular', kind=ChartKind.BAR,
                     labels=['A', 'B'], series=[{'name': 'S', 'values': [1]}]),
        ],
    )
    report = assess_document('pptx', spec)
    assert report['warning_count'] >= 2  # empty table + irregular chart
    print(f"  PASS: quality warnings ({report['warning_count']} warnings)")


def test_pptx_overflow_splits_slides():
    spec = DocumentSpec(
        title='Overflow',
        blocks=[
            SpecBlock(type=BlockType.HEADING, level=1, text='Long list'),
            SpecBlock(type=BlockType.BULLETS, items=[f'Item {i}' for i in range(40)]),
        ],
    )
    out = '/tmp/test_overflow_split.pptx'
    r = TOOL_REGISTRY['generate_pptx']({'path': out, 'document_spec': spec.model_dump()})
    prs = Presentation(r['path'])
    assert len(prs.slides) > 1
    assert r['quality']['check_count'] >= 1
    print(f"  PASS: PPTX overflow split ({len(prs.slides)} slides)")


def test_academic_defense_contract_and_notes():
    from PIL import Image

    img_path = '/tmp/test_defense_source_figure.png'
    Image.new('RGB', (640, 360), '#EEF5FF').save(img_path)

    source = SourceRef(document='thesis.pdf', page=12, section='2.1')
    spec = DocumentSpec(
        title='Thesis Defense',
        subtitle='Complex network dismantling',
        author='Himalaya',
        document_type=DocumentType.DEGREE_DEFENSE,
        audience='degree committee',
        source_documents=[source],
        generation_contract=GenerationContract(
            expected_slide_count=4,
            strict_source_grounding=True,
            require_speaker_notes=True,
            required_sections=['Background', 'Method', 'Results', 'Conclusion'],
            required_assets=AssetPlan(figures=1, tables=1, formulas=1, charts=1, require_source_refs=True, require_extracted_assets=True),
        ),
        blocks=[
            SpecBlock(type=BlockType.HEADING, level=1, text='Background',
                      key_message='Research gap is grounded in the source thesis.',
                      speaker_notes='Explain the motivation and committee-facing context.',
                      source_refs=[source]),
            SpecBlock(type=BlockType.IMAGE, path=img_path, caption='Figure 1. Source topology',
                      source_path='thesis.pdf', crop_box=[10, 20, 400, 260], asset_role='figure',
                      source_refs=[SourceRef(document='thesis.pdf', page=18, figure='Fig. 2-1')]),
            SpecBlock(type=BlockType.HEADING, level=1, text='Method', speaker_notes='Walk through the model.',
                      source_refs=[SourceRef(document='thesis.pdf', page=31)]),
            SpecBlock(type=BlockType.FORMULA, latex=r'R=\frac{1}{N}\sum_i r_i',
                      equation_number='Eq. (3-1)',
                      source_path='thesis.pdf', crop_box=[80, 120, 420, 165],
                      source_refs=[SourceRef(document='thesis.pdf', page=35, equation='3-1')]),
            SpecBlock(type=BlockType.HEADING, level=1, text='Results',
                      source_refs=[SourceRef(document='thesis.pdf', page=61)]),
            SpecBlock(type=BlockType.TABLE, caption='Experiment settings',
                      headers=['Parameter', 'Value'], rows=[['N', '1000']],
                      source_path='thesis.pdf', crop_box=[55, 170, 460, 250],
                      source_refs=[SourceRef(document='thesis.pdf', page=58, table='Table 4-1')]),
            SpecBlock(type=BlockType.CHART, title='Robustness comparison', kind=ChartKind.LINE,
                      labels=['0%', '20%'], series=[{'name': 'Proposed', 'values': [1.0, 0.72]}],
                      source_refs=[SourceRef(document='thesis.pdf', page=65, figure='Fig. 4-3')]),
            SpecBlock(type=BlockType.HEADING, level=1, text='Conclusion',
                      speaker_notes='Close the research loop.',
                      source_refs=[SourceRef(document='thesis.pdf', page=89)]),
            SpecBlock(type=BlockType.BULLETS, items=['Modeling, attack strategy, and validation form a closed loop.'],
                      source_refs=[SourceRef(document='thesis.pdf', page=90)]),
        ],
    )
    out = '/tmp/test_academic_defense.pptx'
    asset_manifest = {
        'complete': True,
        'asset_count': 3,
        'assets': [
            {'role': 'figure', 'path': img_path, 'source_path': 'thesis.pdf', 'page': 18},
            {'role': 'table_image', 'path': img_path, 'source_path': 'thesis.pdf', 'page': 58},
            {'role': 'formula_candidate', 'path': img_path, 'source_path': 'thesis.pdf', 'page': 35},
        ],
        'asset_inventory': [],
    }
    r = TOOL_REGISTRY['generate_pptx']({'path': out, 'document_spec': spec.model_dump(), 'asset_manifest': asset_manifest})
    prs = Presentation(r['path'])
    assert len(prs.slides) == 4
    assert 'Explain the motivation' in prs.slides[0].notes_slide.notes_text_frame.text
    assert r['quality']['source_ref_count'] >= 8
    assert r['quality']['extracted_asset_count'] >= 3
    assert r['quality']['asset_manifest_checked'] is True
    assert r['quality']['quality_level'] != 'failed'
    assert r['quality']['speaker_note_count'] == 3
    manifest = json.loads(open(r['manifestPath'], encoding='utf-8').read())
    assert manifest['document']['document_type'] == 'degree_defense'
    assert manifest['document']['source_documents'][0]['document'] == 'thesis.pdf'
    print("  PASS: academic defense contract, provenance, and notes")


def test_academic_defense_fails_without_extracted_assets_or_formulas():
    source = SourceRef(document='thesis.pdf', page=1)
    spec = DocumentSpec(
        title='Fake grounded defense',
        document_type=DocumentType.DEGREE_DEFENSE,
        source_documents=[source],
        generation_contract=GenerationContract(
            expected_slide_count=12,
            strict_source_grounding=True,
            required_sections=['Background', 'Method', 'Results', 'Conclusion'],
            required_assets=AssetPlan(
                figures=1,
                tables=1,
                formulas=1,
                charts=1,
                require_source_refs=True,
                require_extracted_assets=True,
            ),
        ),
        blocks=[
            SpecBlock(type=BlockType.HEADING, level=1, text='Background', source_refs=[source]),
            SpecBlock(
                type=BlockType.IMAGE,
                caption='Placeholder visual without PDF crop trace',
                source_refs=[SourceRef(document='thesis.pdf', page=2, figure='Fig. 1')],
            ),
            SpecBlock(
                type=BlockType.TABLE,
                caption='Synthetic table',
                headers=['Metric', 'Value'],
                rows=[['Accuracy', '0.91']],
                source_refs=[SourceRef(document='thesis.pdf', page=3, table='Table 1')],
            ),
            SpecBlock(
                type=BlockType.CHART,
                title='Synthetic trend',
                kind=ChartKind.LINE,
                labels=['A', 'B'],
                series=[{'name': 'S', 'values': [1.0, 0.8]}],
                source_refs=[SourceRef(document='thesis.pdf', page=4, figure='Fig. 2')],
            ),
            SpecBlock(type=BlockType.HEADING, level=1, text='Method'),
            SpecBlock(type=BlockType.HEADING, level=1, text='Results'),
            SpecBlock(type=BlockType.HEADING, level=1, text='Conclusion'),
        ],
    )
    report = assess_document('pptx', spec, slide_count=12)
    failures = '\n'.join(report['failures'])
    assert report['extracted_asset_count'] == 0
    assert 'assets.formulas_missing' in failures
    assert 'assets.extracted_missing' in failures
    assert 'assets.manifest_missing' in failures
    assert 'defense.no_formulas' in failures
    assert 'defense.no_extracted_source_assets' in failures
    assert report['quality_level'] == 'failed'
    print(f"  PASS: ungrounded defense visuals rejected ({report['failure_count']} failures)")


def test_formula_candidate_images_count_as_formula_evidence():
    from PIL import Image

    formula_img = '/tmp/test_formula_candidate_asset.png'
    Image.new('RGB', (640, 120), '#FFFFFF').save(formula_img)
    source = SourceRef(document='thesis.pdf', page=21)
    spec = DocumentSpec(
        title='Formula image evidence',
        document_type=DocumentType.DEGREE_DEFENSE,
        source_documents=[source],
        generation_contract=GenerationContract(
            expected_slide_count=4,
            strict_source_grounding=True,
            required_sections=['Background', 'Method', 'Results', 'Conclusion'],
            required_assets=AssetPlan(formulas=1, require_source_refs=True, require_extracted_assets=True),
        ),
        blocks=[
            SpecBlock(type=BlockType.HEADING, level=1, text='Background', source_refs=[source]),
            SpecBlock(type=BlockType.HEADING, level=1, text='Method', source_refs=[source]),
            SpecBlock(
                type=BlockType.IMAGE,
                path=formula_img,
                asset_role='formula_candidate',
                source_path='thesis.pdf',
                crop_box=[40, 100, 500, 145],
                source_refs=[SourceRef(document='thesis.pdf', page=21, equation='2-1')],
            ),
            SpecBlock(type=BlockType.HEADING, level=1, text='Results', source_refs=[source]),
            SpecBlock(type=BlockType.HEADING, level=1, text='Conclusion', source_refs=[source]),
        ],
    )
    report = assess_document('pptx', spec, slide_count=4)
    joined = '\n'.join(report['failures'])
    assert report['formula_image_count'] == 1
    assert 'assets.formulas_missing' not in joined
    assert 'defense.no_formulas' not in joined
    print("  PASS: formula candidate images count as formula evidence")


def test_academic_failure_pattern_is_flagged():
    spec = DocumentSpec(
        document_type=DocumentType.DEGREE_DEFENSE,
        generation_contract=GenerationContract(
            expected_slide_count=18,
            strict_source_grounding=True,
            required_assets=AssetPlan(figures=2, tables=1, formulas=1, charts=1, require_source_refs=True),
            required_sections=['Background', 'Method', 'Results', 'Conclusion'],
        ),
        blocks=[
            SpecBlock(type=BlockType.HEADING, level=1, text='[Insert Paper Title]'),
            SpecBlock(type=BlockType.BULLETS, items=['The Problem', 'Observation: As seen in Figure X']),
        ],
    )
    report = assess_document('pptx', spec, slide_count=1)
    warnings = '\n'.join(report['warnings'])
    failures = '\n'.join(report['failures'])
    assert 'content.placeholders' in warnings
    assert 'sources.missing' in failures
    assert 'defense.too_short' in warnings
    assert 'assets.formulas_missing' in failures
    assert report['quality_level'] == 'failed'
    print(f"  PASS: academic failure pattern flagged ({report['failure_count']} failures, {report['warning_count']} warnings)")


if __name__ == '__main__':
    print("=== Himalaya Doc Service Comprehensive Tests ===\n")
    for name, fn in list(globals().items()):
        if name.startswith('test_') and callable(fn):
            fn()
    print(f"\n=== ALL TESTS PASSED ===")
